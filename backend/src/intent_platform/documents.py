from io import BytesIO
from pathlib import Path
from zipfile import ZipFile


def extract_document(filename, data):
    ext = Path(filename).suffix.lower()
    if ext in {".txt", ".md", ".json", ".yaml", ".yml"}:
        return data.decode("utf-8-sig")
    if ext == ".docx":
        from docx import Document
        with ZipFile(BytesIO(data)) as z:
            if sum(x.file_size for x in z.infolist()) > 50_000_000 or len(z.infolist()) > 5000:
                raise ValueError("DOCX expanded content exceeds limits")
        doc = Document(BytesIO(data))
        return "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted or len(reader.pages) > 300:
            raise ValueError("Encrypted PDFs or documents over 300 pages are not supported")
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
        if not text.strip():
            raise ValueError("Scanned PDF has no extractable text; supply readable text")
        return text
    raise ValueError("Supported formats: TXT, Markdown, DOCX, text PDF")
